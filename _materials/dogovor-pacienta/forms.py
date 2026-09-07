#!/usr/bin/env python3
"""Бланки, которые выдаются вместе с договором: ИДС и медицинская карта 043/у.

    python3 _materials/dogovor-pacienta/forms.py          # два .docx
    python3 _materials/dogovor-pacienta/forms.py --pdf    # + pdf для просмотра

Тексты — в `forms_text.py`, вёрстка использует те же помощники, что и договор
(`build.py`): поля, таблицы-формы, колонтитул. Так три документа выглядят
одинаково и правятся в одном месте.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402  — помощники вёрстки и сборка pdf
import contract_text as C  # noqa: E402
import forms_text as F  # noqa: E402

OUT = HERE / "out"
IDS_NAME = "Информированное-добровольное-согласие"
CARD_NAME = "Медицинская-карта-043у"


def clinic_head(doc):
    """Шапка бланка: клиника и лицензия — заполнять не нужно."""
    c = C.COMPANY
    B.p(doc, c["full"], bold=True, align="left", space=1, size=9.5)
    B.p(doc, f"Место оказания услуг: стоматология {C.CLINIC['name']}, "
             f"{C.CLINIC['address']}, тел. {C.CLINIC['phone']}",
        align="left", space=1, size=9.5)
    B.p(doc, f"Лицензия на осуществление медицинской деятельности "
             f"№ {c['license']}, предоставлена бессрочно",
        align="left", space=6, size=9.5)


def title(doc, text, sub=None):
    for line in text.split("\n"):
        B.p(doc, line, bold=True, align="center", size=12.5, space=2,
            keep_with_next=True)
    if sub:
        B.p(doc, sub, align="center", size=9, space=8, italic=True)
    else:
        B.p(doc, "", space=6)


def checklist(doc, rows):
    """Вмешательства с колонкой для отметки: печатный текст + пустая клетка."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    B.set_widths(table, [14.0, 3.0])
    head = table.rows[0].cells
    for cell, text in zip(head, ("Вмешательство и возможные осложнения",
                                 "Согласен(на)")):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        B._set_font(par.add_run(text), 9, bold=True)
    for row, (name, risks) in zip(table.rows[1:], rows):
        row.cant_split = True
        cell = row.cells[0]
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.space_after = Pt(0)
        B._set_font(par.add_run(name + ". "), 9, bold=True)
        B._set_font(par.add_run("Возможные осложнения: " + risks), 9)
        mark = row.cells[1]
        mark.text = ""
        mark.paragraphs[0].paragraph_format.space_after = Pt(0)
    return table


def build_ids() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Информированное добровольное согласие на медицинское "
                      "вмешательство")
    clinic_head(doc)
    title(doc, F.IDS_TITLE, F.IDS_INTRO)

    B.p(doc, "Часть 1. Согласие на виды вмешательств, включённые в Перечень "
             "(форма по приказу Минздрава России от 12.11.2021 № 1051н)",
        bold=True, align="left", size=10.5, space=3, keep_with_next=True)
    B.form_table(doc, F.IDS_PERSON_ROWS,
                 tall={r[0] for r in F.IDS_PERSON_ROWS[:1]})
    B.p(doc, "", space=3)
    B.p(doc, "Заполняется, если согласие подписывает законный представитель:",
        align="left", size=9, space=2, italic=True)
    B.form_table(doc, F.IDS_CHILD_ROWS)
    B.p(doc, "", space=4)
    for block in F.IDS_BODY:
        B.p(doc, block, space=3)

    doc.add_page_break()
    B.p(doc, "Часть 2. Согласие на стоматологические вмешательства",
        bold=True, align="left", size=10.5, space=2, keep_with_next=True)
    B.p(doc, "Врач отмечает вмешательства, на которые даётся согласие; риски и "
             "возможные осложнения по каждому из них приведены здесь же.",
        align="left", size=9, space=4, italic=True)
    checklist(doc, F.IDS_PROCEDURES)
    B.p(doc, "", space=5)
    for block in F.IDS_EXPLAIN:
        B.p(doc, block, space=2)

    B.p(doc, "", space=4)
    B.p(doc, F.IDS_DISCLOSURE, space=3, keep_with_next=True)
    B.form_table(doc, [("Фамилия, имя, отчество, контактный телефон", ""),
                       ("Фамилия, имя, отчество, контактный телефон", "")])
    B.p(doc, "", space=6)
    B.form_table(doc, F.IDS_SIGN_ROWS,
                 tall={F.IDS_SIGN_ROWS[0][0], F.IDS_SIGN_ROWS[1][0]})

    path = OUT / f"{IDS_NAME}.docx"
    doc.save(path)
    return path


def teeth_table(doc, top, bottom):
    """Зубная формула: два ряда номеров и по ряду клеток под записи."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    table = doc.add_table(rows=4, cols=16)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    B.set_widths(table, [17.0 / 16] * 16)
    for r, values in ((0, top), (3, bottom)):
        for cell, num in zip(table.rows[r].cells, values):
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            B._set_font(par.add_run(num), 8)
    for r in (1, 2):
        table.rows[r].height = Cm(0.75)
        for cell in table.rows[r].cells:
            cell.text = ""
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return table


def build_card() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Медицинская карта стоматологического пациента, форма № 043/у")
    clinic_head(doc)
    title(doc, F.CARD_TITLE, F.CARD_SUB)

    B.form_table(doc, F.CARD_HEAD_ROWS,
                 tall={"Фамилия, имя, отчество", "Адрес места жительства, телефон",
                       "Документ, удостоверяющий личность",
                       "Законный представитель (для детей до 15 лет)"})
    B.p(doc, "", space=8)

    B.p(doc, "Анкета о состоянии здоровья", bold=True, align="left", size=10.5,
        space=2, keep_with_next=True)
    B.p(doc, "Заполняется при первичном обращении, уточняется при каждом "
             "посещении. Отметьте «да» или «нет»; при ответе «да» — уточните.",
        align="left", size=9, space=4, italic=True)
    B.p(doc, F.CARD_ANAMNESIS_NOTE, size=8.5, space=3, italic=True)
    B.form_table(doc, [(q, "да / нет") for q in F.CARD_ANAMNESIS] +
                      [("Достоверность сведений подтверждаю. Подпись Пациента "
                        "(законного представителя), дата", "")],
                 label_w=11.8, row_h=0.52,
                 tall={"Достоверность сведений подтверждаю. Подпись Пациента "
                       "(законного представителя), дата"})
    doc.add_page_break()

    B.p(doc, "Зубная формула", bold=True, align="left", size=10.5, space=3,
        keep_with_next=True)
    teeth_table(doc, F.CARD_TEETH_TOP, F.CARD_TEETH_BOTTOM)
    B.p(doc, "", space=2)
    B.p(doc, F.CARD_TEETH_NOTE, size=8.5, space=8, italic=True)

    for name, lines in F.CARD_EXAM_BLOCKS:
        B.p(doc, name, bold=True, align="left", size=9.5, space=1,
            keep_with_next=True)
        block = doc.add_table(rows=1, cols=1)
        block.style = "Table Grid"
        B.set_widths(block, [17.0])
        block.rows[0].height = Cm(0.55 * lines)
        block.rows[0].cells[0].text = ""
        block.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
        B.p(doc, "", space=3)

    B.p(doc, "", space=6)
    B.p(doc, "Дневник приёмов", bold=True, align="left", size=10.5, space=3,
        keep_with_next=True)
    diary = doc.add_table(rows=F.CARD_DIARY_ROWS + 1, cols=3)
    diary.style = "Table Grid"
    B.set_widths(diary, [2.4, 10.6, 4.0])
    for cell, head in zip(diary.rows[0].cells, F.CARD_DIARY_HEAD):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        B._set_font(par.add_run(head), 9, bold=True)
    for row in diary.rows[1:]:
        row.height = Cm(1.5)
        for cell in row.cells:
            cell.text = ""
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    B.p(doc, "", space=3)
    B.p(doc, F.CARD_DIARY_NOTE, size=9, italic=True)

    path = OUT / f"{CARD_NAME}.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_ids(), build_card()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
