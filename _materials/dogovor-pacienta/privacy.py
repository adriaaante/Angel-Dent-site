#!/usr/bin/env python3
"""Документы по персональным данным → .docx (по файлу на документ).

    python3 _materials/dogovor-pacienta/privacy.py          # все документы
    python3 _materials/dogovor-pacienta/privacy.py --pdf    # + pdf

Тексты — в `privacy_text.py`; вёрстка — общие помощники из `build.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402
import consents as K  # noqa: E402  — bullets/heading
import contract_text as C  # noqa: E402
import forms as F  # noqa: E402
import office as O  # noqa: E402  — approval_block
import papers as P  # noqa: E402  — grid
import privacy_text as X  # noqa: E402

OUT = HERE / "out"


def build_policy() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Политика в отношении обработки персональных данных")
    F.clinic_head(doc)
    O.approval_block(doc)
    F.title(doc, X.POLICY_TITLE, X.POLICY_INTRO)
    for i, (name, items) in enumerate(X.POLICY):
        K.heading(doc, name)
        K.bullets(doc, items)
        if i < len(X.POLICY) - 1:      # хвостовой пустой абзац уводил подписи
            B.p(doc, "", space=4)      # на отдельную пустую страницу
    path = OUT / "Политика-обработки-персональных-данных.docx"
    doc.save(path)
    return path


def build_order() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Приказ об организации обработки и защиты персональных данных")
    F.clinic_head(doc)
    B.p(doc, "", space=6)
    B.p(doc, X.ORDER_TITLE, bold=True, align="center", size=13, space=2,
        keep_with_next=True)
    B.p(doc, X.ORDER_SUB, align="center", size=10.5, space=8, keep_with_next=True)

    head = B.borderless(doc.add_table(rows=1, cols=2))
    head.autofit = True
    B.cell_text(head.rows[0].cells[0], [C.CLINIC["city"]])
    par = head.rows[0].cells[1].paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    B._set_font(par.add_run("№ _______ от «____» ______________ 20___ г."))
    B.p(doc, "", space=6)

    for block in X.ORDER_BODY:
        B.p(doc, block, first_line=0.75, space=4)
    B.p(doc, "ПРИКАЗЫВАЮ:", bold=True, align="left", space=4, keep_with_next=True)
    for i, item in enumerate(X.ORDER_ITEMS, start=1):
        B.clause(doc, f"{i}.", item)
    B.p(doc, "", space=10)
    B.form_table(doc, [
        (C.SIGNATORY["short_position"],
         f"__________________ / {C.SIGNATORY['short_name']}"),
        ("С приказом ознакомлены (Ф. И. О., подпись, дата)", ""),
    ], tall={C.SIGNATORY["short_position"]},
       extra_tall={"С приказом ознакомлены (Ф. И. О., подпись, дата)"})

    # Приложение № 2 — перечень допущенных лиц (приложение № 1 — сама Политика)
    doc.add_page_break()
    F.clinic_head(doc)
    F.title(doc, X.ACCESS_TITLE, X.ACCESS_INTRO)
    P.grid(doc, X.ACCESS_HEAD, [1.0, 4.6, 3.6, 5.0, 2.8], X.ACCESS_ROWS, row_h=0.7)
    B.p(doc, "", space=8)
    B.form_table(doc, [
        ("Ответственный за организацию обработки персональных данных", ""),
    ], tall={"Ответственный за организацию обработки персональных данных"})
    path = OUT / "Приказ-об-организации-обработки-персональных-данных.docx"
    doc.save(path)
    return path


def build_nda() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Обязательство о неразглашении врачебной тайны и "
                      "персональных данных")
    F.clinic_head(doc)
    F.title(doc, X.NDA_TITLE, X.NDA_INTRO)

    B.form_table(doc, X.NDA_ROWS, tall={X.NDA_ROWS[0][0]}, row_h=0.5)
    B.p(doc, "", space=4)
    for block in X.NDA_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=3)
    B.p(doc, "Принимаю на себя обязательства:", bold=True, space=3,
        keep_with_next=True)
    K.bullets(doc, X.NDA_ITEMS, size=10)
    B.p(doc, "", space=3)
    B.p(doc, X.NDA_TAIL, space=5)
    B.form_table(doc, X.NDA_SIGN_ROWS, tall={X.NDA_SIGN_ROWS[0][0]}, row_h=0.5)
    path = OUT / "Обязательство-о-неразглашении-врачебной-тайны.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_policy(), build_order(), build_nda()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
