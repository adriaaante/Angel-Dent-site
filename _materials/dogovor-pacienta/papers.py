#!/usr/bin/env python3
"""Рабочие документы по лечению и оплате → .docx (по файлу на документ).

    python3 _materials/dogovor-pacienta/papers.py          # все бланки
    python3 _materials/dogovor-pacienta/papers.py --pdf    # + pdf

Тексты — в `papers_text.py`; вёрстка — общие помощники из `build.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402
import consents as K  # noqa: E402  — bullets/heading/checkboxes
import forms as F  # noqa: E402
import papers_text as X  # noqa: E402

OUT = HERE / "out"


def grid(doc, head, widths, rows, row_h=0.62, size=9.5):
    """Таблица с печатной шапкой и пустыми строками под заполнение."""
    table = doc.add_table(rows=rows + 1, cols=len(head))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    B.set_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, head):
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        B._set_font(par.add_run(text), size, bold=True)
    for row in table.rows[1:]:
        row.height = Cm(row_h)
        for cell in row.cells:
            cell.text = ""
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return table


def build_plan() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "План лечения (смета) — приложение № 1 к договору")
    F.clinic_head(doc)
    F.title(doc, X.PLAN_TITLE, X.PLAN_INTRO)

    B.form_table(doc, X.PLAN_HEAD_ROWS,
                 tall={X.PLAN_HEAD_ROWS[0][0], X.PLAN_HEAD_ROWS[3][0]},
                 extra_tall={X.PLAN_HEAD_ROWS[4][0]})
    B.p(doc, "", space=6)

    K.heading(doc, "Перечень услуг и стоимость")
    grid(doc, X.PLAN_TABLE_HEAD, [1.0, 9.0, 1.8, 2.6, 2.6], X.PLAN_ROWS)
    B.p(doc, "", space=6)

    K.heading(doc, "Этапы и сроки")
    grid(doc, X.PLAN_STAGES_HEAD, [1.6, 8.0, 4.0, 3.4], X.PLAN_STAGES_ROWS, row_h=0.8)
    B.p(doc, "", space=6)

    K.heading(doc, "Гарантийные сроки и сроки службы")
    grid(doc, X.PLAN_WARRANTY_HEAD, [8.0, 4.5, 4.5], X.PLAN_WARRANTY_ROWS)
    B.p(doc, "", space=6)

    B.p(doc, X.PLAN_TYPE, space=5)
    for note in X.PLAN_NOTES:
        B.p(doc, note, space=3)
    B.p(doc, "", space=5)
    B.form_table(doc, X.PLAN_SIGN_ROWS,
                 tall={X.PLAN_SIGN_ROWS[0][0], X.PLAN_SIGN_ROWS[1][0],
                       X.PLAN_SIGN_ROWS[3][0]},
                 extra_tall={X.PLAN_SIGN_ROWS[2][0]})
    path = OUT / "Приложение-1-План-лечения-смета.docx"
    doc.save(path)
    return path


def build_addendum() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Дополнительное соглашение об изменении плана лечения")
    F.clinic_head(doc)
    F.title(doc, X.ADD_TITLE, X.ADD_INTRO)

    B.form_table(doc, X.ADD_HEAD_ROWS, tall={X.ADD_HEAD_ROWS[0][0]})
    B.p(doc, "", space=5)
    for block in X.ADD_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=4)
    B.form_table(doc, X.ADD_REASON_ROWS, extra_tall={X.ADD_REASON_ROWS[0][0]})
    B.p(doc, "", space=5)

    K.heading(doc, "Изменения в плане лечения")
    grid(doc, X.ADD_TABLE_HEAD, [1.0, 4.6, 6.4, 1.8, 3.2], X.ADD_ROWS)
    B.p(doc, "", space=5)
    B.form_table(doc, X.ADD_TOTAL_ROWS, tall={X.ADD_TOTAL_ROWS[1][0]})
    B.p(doc, "", space=6)
    B.form_table(doc, X.ADD_SIGN_ROWS, extra_tall={X.ADD_SIGN_ROWS[0][0]},
                 tall={X.ADD_SIGN_ROWS[1][0]})
    path = OUT / "Дополнительное-соглашение-об-изменении-плана-лечения.docx"
    doc.save(path)
    return path


def build_instalment() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Соглашение о рассрочке оплаты")
    F.clinic_head(doc)
    F.title(doc, X.INSTALMENT_TITLE, X.INSTALMENT_INTRO)

    B.form_table(doc, X.INSTALMENT_HEAD_ROWS,
                 tall={X.INSTALMENT_HEAD_ROWS[0][0], X.INSTALMENT_HEAD_ROWS[1][0],
                       X.INSTALMENT_HEAD_ROWS[3][0]})
    B.p(doc, "", space=6)
    K.heading(doc, "График платежей")
    grid(doc, X.INSTALMENT_TABLE_HEAD, [2.6, 4.0, 5.4, 5.0], X.INSTALMENT_ROWS)
    B.p(doc, "", space=6)
    for block in X.INSTALMENT_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=6)
    B.form_table(doc, X.INSTALMENT_SIGN_ROWS,
                 tall=set(r[0] for r in X.INSTALMENT_SIGN_ROWS))
    path = OUT / "Соглашение-о-рассрочке-оплаты.docx"
    doc.save(path)
    return path


def build_act() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Акт об оказании медицинских услуг и гарантийный талон")
    F.clinic_head(doc)
    F.title(doc, X.ACT_TITLE, X.ACT_INTRO)

    B.form_table(doc, X.ACT_HEAD_ROWS,
                 tall={X.ACT_HEAD_ROWS[0][0], X.ACT_HEAD_ROWS[1][0],
                       X.ACT_HEAD_ROWS[2][0]})
    B.p(doc, "", space=6)
    K.heading(doc, "Оказанные услуги")
    grid(doc, X.ACT_TABLE_HEAD, [1.0, 8.2, 2.4, 1.8, 3.0], X.ACT_ROWS)
    B.p(doc, "", space=5)
    B.form_table(doc, X.ACT_TOTAL_ROWS, tall={X.ACT_TOTAL_ROWS[0][0]})
    B.p(doc, "", space=5)
    for block in X.ACT_BODY:
        B.p(doc, block, space=3)

    B.p(doc, "", space=6)
    K.heading(doc, X.ACT_WARRANTY_TITLE)
    grid(doc, X.ACT_WARRANTY_HEAD, [5.6, 4.6, 3.4, 3.4], X.ACT_WARRANTY_ROWS)
    B.p(doc, "", space=4)
    for note in X.ACT_WARRANTY_NOTES:
        B.p(doc, note, space=3, size=9.5)
    B.p(doc, "", space=6)
    B.form_table(doc, X.ACT_SIGN_ROWS, extra_tall={X.ACT_SIGN_ROWS[0][0]},
                 tall={X.ACT_SIGN_ROWS[1][0]})
    path = OUT / "Акт-об-оказании-услуг-и-гарантийный-талон.docx"
    doc.save(path)
    return path


def build_memos() -> Path:
    """Памятки: каждая на своей странице — врач печатает нужную."""
    doc = B.docx_base()
    B.add_footer(doc, "Памятки пациенту: рекомендации после лечения")
    F.clinic_head(doc)
    F.title(doc, X.MEMO_TITLE, X.MEMO_INTRO)

    for i, (name, items) in enumerate(X.MEMOS):
        if i:
            doc.add_page_break()
            F.clinic_head(doc)
        K.heading(doc, name)
        K.bullets(doc, items)
        B.p(doc, "", space=5)
        B.form_table(doc, X.MEMO_SIGN_ROWS,
                     tall=set(r[0] for r in X.MEMO_SIGN_ROWS))
    path = OUT / "Памятки-пациенту-после-лечения.docx"
    doc.save(path)
    return path


def build_claim() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Заявление об отказе от исполнения договора и возврате средств")

    for line in X.CLAIM_ADDRESS.split("\n"):
        B.p(doc, line, align="right", space=1)
    B.p(doc, "", space=6)
    F.title(doc, X.CLAIM_TITLE, X.CLAIM_INTRO)

    B.form_table(doc, X.CLAIM_ROWS,
                 tall={X.CLAIM_ROWS[0][0], X.CLAIM_ROWS[1][0], X.CLAIM_ROWS[3][0],
                       X.CLAIM_ROWS[4][0]})
    B.p(doc, "", space=5)
    for block in X.CLAIM_BODY:
        B.p(doc, block, space=3)
    B.form_table(doc, X.CLAIM_MONEY_ROWS,
                 tall={X.CLAIM_MONEY_ROWS[0][0]},
                 extra_tall={X.CLAIM_MONEY_ROWS[1][0], X.CLAIM_MONEY_ROWS[2][0]})
    B.p(doc, "", space=4)
    for block in X.CLAIM_TAIL:
        B.p(doc, block, space=3)
    B.p(doc, "", space=5)
    B.form_table(doc, X.CLAIM_SIGN_ROWS, tall={X.CLAIM_SIGN_ROWS[0][0]})

    B.p(doc, "", space=8)
    K.heading(doc, X.CLAIM_OFFICE_TITLE)
    B.form_table(doc, X.CLAIM_OFFICE_ROWS,
                 tall={X.CLAIM_OFFICE_ROWS[0][0], X.CLAIM_OFFICE_ROWS[3][0]},
                 extra_tall={X.CLAIM_OFFICE_ROWS[1][0], X.CLAIM_OFFICE_ROWS[2][0]})
    path = OUT / "Заявление-о-возврате-денежных-средств.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_plan(), build_addendum(), build_instalment(), build_act(),
            build_memos(), build_claim()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
