#!/usr/bin/env python3
"""Согласия и отказы пациента → .docx (по одному файлу на документ).

    python3 _materials/dogovor-pacienta/consents.py          # все бланки
    python3 _materials/dogovor-pacienta/consents.py --pdf    # + pdf

Тексты — в `consents_text.py`, вёрстка — те же помощники, что у договора
(`build.py`) и у бланков карты (`forms.py`): один вид у всего пакета.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402
import consents_text as X  # noqa: E402
import forms as F  # noqa: E402

OUT = HERE / "out"


def bullets(doc, items, size=10.5):
    """Маркированный список с висячим отступом (без стилей Word)."""
    for item in items:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = B.Cm(0.8)
        pf.first_line_indent = B.Cm(-0.4)
        pf.space_after = Pt(2)
        B._set_font(par.add_run("— "), size)
        B._set_font(par.add_run(B.placeholders(item)), size)


def checkboxes(doc, items, width=(4.4, 12.6)):
    """Строки «[ ] Согласен(на) / [ ] Не согласен(на) — описание»."""
    table = B.set_widths(B.borderless(doc.add_table(rows=len(items), cols=2)),
                         list(width))
    for row, text in zip(table.rows, items):
        row.cant_split = True
        B.cell_text(row.cells[0], ["[   ] Согласен(на)", "[   ] Не согласен(на)"])
        B.cell_text(row.cells[1], [text])
    return table


def heading(doc, text):
    B.p(doc, text, bold=True, align="left", size=11, space=3, keep_with_next=True)


def build_consent(item: dict) -> Path:
    """ИДС по виду лечения: цели и методы → риски → особенности → подписи."""
    doc = B.docx_base()
    short = item["title"].split("\n")[-1]
    B.add_footer(doc, f"Информированное добровольное согласие: {short[:70]}")
    F.clinic_head(doc)
    F.title(doc, item["title"], X.IDS_LEAD)

    B.form_table(doc, X.PATIENT_ROWS,
                 tall={X.PATIENT_ROWS[0][0], X.PATIENT_ROWS[2][0],
                       X.PATIENT_ROWS[3][0], X.PATIENT_ROWS[4][0]})
    B.p(doc, "", space=6)

    heading(doc, "Цель, методы и альтернативы")
    for block in item["about"]:
        B.p(doc, block, space=3)

    B.p(doc, "", space=3)
    heading(doc, "Возможные осложнения и последствия, о которых я предупреждён(а)")
    bullets(doc, item["risks"])

    B.p(doc, "", space=3)
    heading(doc, "Особенности этого лечения")
    for block in item["extra"]:
        B.p(doc, block, space=3)

    B.p(doc, "", space=3)
    B.p(doc, X.ANESTHESIA, space=3)
    for block in X.COMMON_TAIL:
        B.p(doc, block, space=3)
    B.p(doc, "Я даю информированное добровольное согласие на указанное медицинское "
             "вмешательство на изложенных условиях.", bold=True, space=6)

    B.form_table(doc, X.SIGN_ROWS, tall={X.SIGN_ROWS[0][0], X.SIGN_ROWS[1][0]})
    path = OUT / f"{item['file']}.docx"
    doc.save(path)
    return path


def build_refusal() -> Path:
    """Отказ от вмешательств из Перечня — форма приказа № 1051н, дословно."""
    doc = B.docx_base()
    B.add_footer(doc, "Отказ от медицинского вмешательства")
    F.clinic_head(doc)
    F.title(doc, X.REFUSAL_TITLE, X.REFUSAL_INTRO)

    B.form_table(doc, X.REFUSAL_ROWS,
                 tall={r[0] for r in X.REFUSAL_ROWS})
    B.p(doc, "", space=4)
    B.p(doc, X.REFUSAL_BODY_1, space=4)
    B.form_table(doc, [X.REFUSAL_FIELDS[0]], tall={X.REFUSAL_FIELDS[0][0]})
    B.p(doc, "", space=4)
    B.p(doc, X.REFUSAL_BODY_2, space=4)
    B.form_table(doc, X.REFUSAL_FIELDS[1:], tall={X.REFUSAL_FIELDS[1][0],
                                                  X.REFUSAL_FIELDS[2][0]})
    B.p(doc, "", space=4)
    B.p(doc, X.REFUSAL_BODY_3, space=6)
    B.form_table(doc, [
        ("Подпись гражданина (законного представителя), расшифровка", ""),
        ("Фамилия, имя, отчество и подпись медицинского работника", ""),
        ("Дата оформления", ""),
    ], tall={"Подпись гражданина (законного представителя), расшифровка",
             "Фамилия, имя, отчество и подпись медицинского работника"})
    path = OUT / "Отказ-от-медицинского-вмешательства.docx"
    doc.save(path)
    return path


def build_partial_refusal() -> Path:
    """Отказ от части рекомендованного лечения — бланк клиники."""
    doc = B.docx_base()
    B.add_footer(doc, "Отказ от рекомендованного объёма лечения")
    F.clinic_head(doc)
    F.title(doc, X.PARTIAL_TITLE, X.PARTIAL_INTRO)

    B.form_table(doc, X.PARTIAL_ROWS,
                 tall={X.PARTIAL_ROWS[0][0], X.PARTIAL_ROWS[2][0]},
                 extra_tall={X.PARTIAL_ROWS[3][0], X.PARTIAL_ROWS[4][0]})
    B.p(doc, "", space=5)
    for block in X.PARTIAL_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=4)
    B.form_table(doc, X.PARTIAL_DOCTOR_ROWS,
                 tall={X.PARTIAL_DOCTOR_ROWS[1][0], X.PARTIAL_DOCTOR_ROWS[2][0]},
                 extra_tall={X.PARTIAL_DOCTOR_ROWS[0][0]})
    path = OUT / "Отказ-от-рекомендованного-объёма-лечения.docx"
    doc.save(path)
    return path


def build_pdn() -> Path:
    """Согласие на обработку персональных данных — отдельным документом."""
    doc = B.docx_base()
    B.add_footer(doc, "Согласие на обработку персональных данных")
    F.clinic_head(doc)
    F.title(doc, X.PDN_TITLE, X.PDN_INTRO)

    B.p(doc, X.PDN_OPERATOR, space=6)
    heading(doc, "Субъект персональных данных")
    B.form_table(doc, X.PDN_SUBJECT_ROWS,
                 tall={X.PDN_SUBJECT_ROWS[0][0], X.PDN_SUBJECT_ROWS[1][0],
                       X.PDN_SUBJECT_ROWS[3][0]},
                 extra_tall={X.PDN_SUBJECT_ROWS[2][0], X.PDN_SUBJECT_ROWS[4][0]})
    B.p(doc, "", space=6)

    heading(doc, "Цели обработки")
    bullets(doc, X.PDN_PURPOSES)
    B.p(doc, "", space=3)
    heading(doc, "Перечень данных")
    B.p(doc, X.PDN_DATA, space=3)
    heading(doc, "Действия с данными")
    B.p(doc, X.PDN_ACTIONS, space=3)
    heading(doc, "Срок действия и отзыв согласия")
    B.p(doc, X.PDN_TERM, space=6)

    heading(doc, "Отметьте, на что даётся согласие")
    checkboxes(doc, X.PDN_CHECKS)
    B.p(doc, "", space=6)
    B.form_table(doc, [
        ("Подпись субъекта персональных данных, расшифровка", ""),
        ("Дата", ""),
    ], tall={"Подпись субъекта персональных данных, расшифровка"})
    path = OUT / "Согласие-на-обработку-персональных-данных.docx"
    doc.save(path)
    return path


def build_photo() -> Path:
    """Согласие на съёмку и публикацию изображений (ст. 152.1 ГК РФ)."""
    doc = B.docx_base()
    B.add_footer(doc, "Согласие на фотографирование, видеосъёмку и использование "
                      "изображений")
    F.clinic_head(doc)
    F.title(doc, X.PHOTO_TITLE, X.PHOTO_INTRO)

    B.form_table(doc, [
        ("Фамилия, имя, отчество", ""),
        ("Дата рождения", ""),
        ("Законный представитель (для Пациента до 18 лет): Ф. И. О., документ о "
         "полномочиях", ""),
    ], tall={"Фамилия, имя, отчество",
             "Законный представитель (для Пациента до 18 лет): Ф. И. О., документ о "
             "полномочиях"})
    B.p(doc, "", space=6)

    for block in X.PHOTO_BLOCKS:
        B.p(doc, block, space=4)
    checkboxes(doc, X.PHOTO_CHECKS)
    B.p(doc, "", space=5)
    heading(doc, "Условия использования")
    for block in X.PHOTO_TERMS:
        B.p(doc, block, space=3)
    B.p(doc, "", space=6)
    B.form_table(doc, [
        ("Подпись Пациента (законного представителя), расшифровка", ""),
        ("Дата", ""),
    ], tall={"Подпись Пациента (законного представителя), расшифровка"})
    path = OUT / "Согласие-на-фото-и-видеосъёмку.docx"
    doc.save(path)
    return path


def build_disclosure() -> Path:
    """Кому клиника вправе сообщать сведения о лечении."""
    doc = B.docx_base()
    B.add_footer(doc, "Согласие на передачу сведений, составляющих врачебную тайну")
    F.clinic_head(doc)
    F.title(doc, X.DISCLOSURE_TITLE, X.DISCLOSURE_INTRO)

    B.form_table(doc, X.DISCLOSURE_ROWS, tall={X.DISCLOSURE_ROWS[0][0]})
    B.p(doc, "", space=5)
    B.p(doc, "В соответствии с пунктом 5 части 5 статьи 19 и частью 3 статьи 13 "
             "Федерального закона от 21.11.2011 № 323-ФЗ разрешаю передавать "
             "сведения, составляющие врачебную тайну, следующим лицам:", space=4)
    B.form_table(doc, X.DISCLOSURE_PERSON_ROWS,
                 tall={r[0] for r in X.DISCLOSURE_PERSON_ROWS})
    B.p(doc, "", space=5)
    for block in X.DISCLOSURE_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=6)
    B.form_table(doc, [
        ("Подпись Пациента (законного представителя), расшифровка", ""),
        ("Дата", ""),
    ], tall={"Подпись Пациента (законного представителя), расшифровка"})
    path = OUT / "Согласие-на-информирование-третьих-лиц.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_consent(item) for item in X.CONSENTS]
    made += [build_refusal(), build_partial_refusal(), build_pdn(),
             build_photo(), build_disclosure()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
