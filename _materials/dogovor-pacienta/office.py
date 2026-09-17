#!/usr/bin/env python3
"""Внутренние документы клиники → .docx (правила, положения, приказ, памятка).

    python3 _materials/dogovor-pacienta/office.py          # все документы
    python3 _materials/dogovor-pacienta/office.py --pdf    # + pdf

Тексты — в `office_text.py`; вёрстка — общие помощники из `build.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402
import consents as K  # noqa: E402  — bullets/heading
import contract_text as C  # noqa: E402
import forms as F  # noqa: E402
import office_text as X  # noqa: E402
import papers as P  # noqa: E402  — grid

OUT = HERE / "out"


def approval_block(doc):
    """Гриф утверждения в правом верхнем углу — как в локальных актах."""
    table = B.set_widths(B.borderless(doc.add_table(rows=1, cols=2)), [9.0, 8.0])
    B.cell_text(table.rows[0].cells[0], [""])
    lines = ["УТВЕРЖДАЮ", f"{C.SIGNATORY_ORDERS['short_position']} {C.COMPANY['short']}",
             f"__________________ / {C.SIGNATORY_ORDERS['short_name']}",
             "«____» ______________ 20___ г."]
    cell = table.rows[0].cells[1]
    cell.text = ""
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.space_after = Pt(1)
        B._set_font(par.add_run(line), 9.5, bold=(i == 0))
    B.p(doc, "", space=6)


def numbered_sections(doc, blocks):
    for name, items in blocks:
        K.heading(doc, name)
        K.bullets(doc, items)
        B.p(doc, "", space=4)


def build_rules() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Правила внутреннего распорядка для пациентов")
    F.clinic_head(doc)
    approval_block(doc)
    F.title(doc, X.RULES_TITLE, X.RULES_INTRO)
    numbered_sections(doc, X.RULES)
    path = OUT / "Правила-внутреннего-распорядка-для-пациентов.docx"
    doc.save(path)
    return path


def build_warranty() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Положение о гарантийных сроках и сроках службы")
    F.clinic_head(doc)
    approval_block(doc)
    F.title(doc, X.WARRANTY_TITLE, X.WARRANTY_INTRO)

    table = P.grid(doc, X.WARRANTY_HEAD, [9.0, 4.0, 4.0], len(X.WARRANTY_ROWS))
    for row, values in zip(table.rows[1:], X.WARRANTY_ROWS):
        for cell, value in zip(row.cells, values):
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            B._set_font(par.add_run(value), 9.5)
    B.p(doc, "", space=6)
    K.heading(doc, "Условия действия гарантии")
    for term in X.WARRANTY_TERMS:
        B.p(doc, term, space=3)
    path = OUT / "Положение-о-гарантийных-сроках.docx"
    doc.save(path)
    return path


def build_cctv() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Положение о системе видеонаблюдения")
    F.clinic_head(doc)
    approval_block(doc)
    F.title(doc, X.CCTV_TITLE, X.CCTV_INTRO)
    numbered_sections(doc, X.CCTV)
    K.heading(doc, "Информирование посетителей")
    B.p(doc, X.CCTV_SIGN, space=3)
    path = OUT / "Положение-о-видеонаблюдении.docx"
    doc.save(path)
    return path


def build_order() -> Path:
    doc = B.docx_base()
    B.add_footer(doc, "Приказ об утверждении форм документов и прейскуранта")
    F.clinic_head(doc)
    B.p(doc, "", space=6)
    B.p(doc, X.ORDER_TITLE, bold=True, align="center", size=13, space=2,
        keep_with_next=True)
    B.p(doc, X.ORDER_SUB, align="center", size=10.5, space=8, keep_with_next=True)

    head = B.borderless(doc.add_table(rows=1, cols=2))
    head.autofit = True
    B.cell_text(head.rows[0].cells[0], [f"{C.CLINIC['city']}"])
    par = head.rows[0].cells[1].paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    B._set_font(par.add_run("№ _______ от «____» ______________ 20___ г."))
    B.p(doc, "", space=6)

    for block in X.ORDER_BODY:
        B.p(doc, block, first_line=0.75, space=4)
    B.p(doc, "ПРИКАЗЫВАЮ:", bold=True, align="left", space=4, keep_with_next=True)
    for i, item in enumerate(X.ORDER_ITEMS, start=1):
        B.clause(doc, f"{i}.", item)
    B.p(doc, "", space=10)

    B.form_table(doc, [
        (f"{C.SIGNATORY_ORDERS['short_position']}",
         f"__________________ / {C.SIGNATORY_ORDERS['short_name']}"),
        ("С приказом ознакомлены (Ф. И. О., подпись, дата)", ""),
    ], tall={f"{C.SIGNATORY_ORDERS['short_position']}"},
       extra_tall={"С приказом ознакомлены (Ф. И. О., подпись, дата)"})
    path = OUT / "Приказ-об-утверждении-документов-и-прейскуранта.docx"
    doc.save(path)
    return path


def build_admin() -> Path:
    doc = B.docx_base()
    for s in doc.sections:            # памятка-таблица — альбомная ориентация
        s.orientation = 1             # WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Cm(1.4)
        s.top_margin = Cm(1.4)
        s.bottom_margin = Cm(1.2)
    B.add_footer(doc, "Памятка администратору и врачу: какой документ когда "
                      "подписывается")
    F.clinic_head(doc, compact=True)
    F.title(doc, X.ADMIN_TITLE, X.ADMIN_INTRO)

    table = P.grid(doc, X.ADMIN_HEAD, [3.8, 6.4, 4.0, 10.0], len(X.ADMIN_ROWS),
                   row_h=0.6)
    for row, values in zip(table.rows[1:], X.ADMIN_ROWS):
        for cell, value in zip(row.cells, values):
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            B._set_font(par.add_run(value), 9)
    B.p(doc, "", space=5)

    # Кегль списков 9,5 и сжатые отступы: памятка должна укладываться в три
    # листа — четвёртый с двумя строками администратор просто не подшивает.
    K.heading(doc, X.ADMIN_MINIMUM_TITLE)
    K.bullets(doc, X.ADMIN_MINIMUM, size=9.5)
    B.p(doc, "", space=4)
    K.heading(doc, X.ADMIN_STAND_TITLE)
    K.bullets(doc, X.ADMIN_STAND, size=9.5)
    B.p(doc, "", space=4)
    K.heading(doc, X.ADMIN_MISTAKES_TITLE)
    K.bullets(doc, X.ADMIN_MISTAKES, size=9.5)
    path = OUT / "Памятка-администратору-какой-документ-когда.docx"
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_rules(), build_warranty(), build_cctv(), build_order(),
            build_admin()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
