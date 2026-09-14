#!/usr/bin/env python3
"""Бланки для нештатных ситуаций → .docx (по файлу на документ).

    python3 _materials/dogovor-pacienta/extra.py          # все бланки
    python3 _materials/dogovor-pacienta/extra.py --pdf    # + pdf

Тексты — в `extra_text.py`; вёрстка — общие помощники из `build.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as B  # noqa: E402
import consents as K  # noqa: E402  — bullets/heading/checkboxes
import extra_text as X  # noqa: E402
import forms as F  # noqa: E402
import papers as P  # noqa: E402  — grid

OUT = HERE / "out"


def box(doc, label, height=2.2):
    """Поле под рукописный текст: подпись сверху, пустая рамка нужной высоты."""
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    B.set_widths(table, [17.0])
    head = table.rows[0].cells[0]
    head.text = ""
    par = head.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    B._set_font(par.add_run(label), 9, bold=True)
    table.rows[0].height = Cm(0.5)
    body = table.rows[1].cells[0]
    body.text = ""
    body.paragraphs[0].paragraph_format.space_after = Pt(0)
    table.rows[1].height = Cm(height)
    B.p(doc, "", space=4)
    return table


def build_inspection() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Акт осмотра результата оказанной медицинской услуги")
    F.clinic_head(doc)
    F.title(doc, X.INSPECT_TITLE, X.INSPECT_INTRO)

    B.form_table(doc, X.INSPECT_HEAD_ROWS,
                 tall={X.INSPECT_HEAD_ROWS[1][0], X.INSPECT_HEAD_ROWS[3][0]},
                 row_h=0.5)
    B.p(doc, "", space=4)
    # Лист заполняется от руки, поэтому места под запись даём щедро: первая
    # страница — что предъявлено и что видит врач, вторая — вывод и решение.
    box(doc, X.INSPECT_COMPLAINT, 3.6)
    box(doc, X.INSPECT_OBJECTIVE, 8.6)

    doc.add_page_break()
    box(doc, X.INSPECT_CONCLUSION, 5.0)
    K.heading(doc, X.INSPECT_DECISION)
    K.checkboxes(doc, X.INSPECT_DECISIONS, width=(3.2, 13.8))
    B.p(doc, "", space=3)
    for note in X.INSPECT_NOTES:
        B.p(doc, note, space=2, size=9)
    B.p(doc, "", space=4)
    B.form_table(doc, X.INSPECT_SIGN_ROWS,
                 extra_tall={r[0] for r in X.INSPECT_SIGN_ROWS})
    path = OUT / "Акт-осмотра-результата-оказанной-услуги.docx"
    doc.save(path)
    return path


def build_reply() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Ответ на обращение (претензию) пациента")
    F.clinic_head(doc)
    F.title(doc, X.REPLY_TITLE, X.REPLY_INTRO)

    B.form_table(doc, X.REPLY_HEAD_ROWS, tall={X.REPLY_HEAD_ROWS[0][0]}, row_h=0.5)
    B.p(doc, "", space=4)
    for block in X.REPLY_BODY:
        B.p(doc, block, space=3)
    B.p(doc, "", space=2)
    K.heading(doc, X.REPLY_OPTIONS_TITLE)
    K.bullets(doc, X.REPLY_OPTIONS, size=9.5)
    B.p(doc, "", space=3)
    for block in X.REPLY_TAIL:
        B.p(doc, block, space=3, size=9.5)
    B.p(doc, "", space=4)
    B.form_table(doc, X.REPLY_SIGN_ROWS, tall={X.REPLY_SIGN_ROWS[0][0]}, row_h=0.5)
    path = OUT / "Ответ-на-обращение-претензию-пациента.docx"
    doc.save(path)
    return path


def build_log() -> Path:
    doc = B.docx_base()
    for s in doc.sections:                 # журнал-таблица — альбомная ориентация
        s.orientation = 1
        s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Cm(1.4)
        s.top_margin = Cm(1.4)
        s.bottom_margin = Cm(1.2)
    B.add_footer(doc, "Журнал учёта обращений, претензий и запросов пациентов")
    F.clinic_head(doc)
    F.title(doc, X.LOG_TITLE, X.LOG_INTRO)
    P.grid(doc, X.LOG_HEAD, [0.9, 2.0, 4.0, 5.2, 5.6, 2.6, 4.6, 2.3], X.LOG_ROWS,
           row_h=0.8, size=8.5)
    B.p(doc, "", space=6)
    K.bullets(doc, X.LOG_NOTES, size=9.5)
    path = OUT / "Журнал-учёта-обращений-и-претензий.docx"
    doc.save(path)
    return path


def build_notice() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Уведомление пациенту")
    F.clinic_head(doc)
    F.title(doc, X.NOTICE_TITLE, X.NOTICE_INTRO)

    B.form_table(doc, X.NOTICE_HEAD_ROWS, tall={X.NOTICE_HEAD_ROWS[0][0],
                                                X.NOTICE_HEAD_ROWS[3][0]})
    B.p(doc, "", space=5)
    K.heading(doc, X.NOTICE_CASES_TITLE)
    table = B.set_widths(B.borderless(doc.add_table(rows=len(X.NOTICE_CASES), cols=2)),
                         [1.6, 15.4])
    for row, text in zip(table.rows, X.NOTICE_CASES):
        row.cant_split = True
        B.cell_text(row.cells[0], ["[   ]"])
        B.cell_text(row.cells[1], [text])
    B.p(doc, "", space=5)
    B.p(doc, X.NOTICE_TAIL, space=6, size=9.5)
    B.form_table(doc, X.NOTICE_SIGN_ROWS,
                 tall={X.NOTICE_SIGN_ROWS[0][0], X.NOTICE_SIGN_ROWS[1][0]},
                 extra_tall={X.NOTICE_SIGN_ROWS[2][0]})
    path = OUT / "Уведомление-пациенту.docx"
    doc.save(path)
    return path


def build_docs_request() -> Path:
    doc = B.docx_base(compact=True)
    B.add_footer(doc, "Заявление о выдаче медицинских документов и справок")
    for line in X.DOCS_ADDRESS.split("\n"):
        B.p(doc, line, align="right", space=1)
    B.p(doc, "", space=6)
    F.title(doc, X.DOCS_TITLE, X.DOCS_INTRO)

    B.form_table(doc, X.DOCS_HEAD_ROWS, row_h=0.5,
                 tall={X.DOCS_HEAD_ROWS[0][0], X.DOCS_HEAD_ROWS[2][0],
                       X.DOCS_HEAD_ROWS[4][0]})
    B.p(doc, "", space=4)
    K.heading(doc, X.DOCS_REQUEST_TITLE)
    table = B.set_widths(B.borderless(doc.add_table(rows=len(X.DOCS_REQUEST), cols=2)),
                         [1.6, 15.4])
    for row, text in zip(table.rows, X.DOCS_REQUEST):
        row.cant_split = True
        B.cell_text(row.cells[0], ["[   ]"])
        B.cell_text(row.cells[1], [text])
    B.p(doc, "", space=3)
    for block in X.DOCS_BODY:
        B.p(doc, block, space=3, size=10)
    B.p(doc, "", space=4)
    B.form_table(doc, X.DOCS_SIGN_ROWS, row_h=0.5)
    B.p(doc, "", space=5)
    K.heading(doc, X.DOCS_OFFICE_TITLE)
    B.p(doc, X.DOCS_NOTE, space=3, size=9)
    B.form_table(doc, X.DOCS_OFFICE_ROWS, row_h=0.55)
    path = OUT / "Заявление-о-выдаче-медицинских-документов.docx"
    doc.save(path)
    return path


# ⚠️ Отдельного бланка «Лист учёта дозовых нагрузок» нет (14.09.2026):
# он печатается на обороте ИДС на рентгенологическое исследование
# (`consents.py` → build_consent, ветка key == "rentgen"). Подписываются они
# в один момент, шапка общая, лист сразу ложится в карту — второй документ
# и вторая шапка были лишними. Тексты листа остались в `extra_text.py`
# (DOSE_*), оттуда их и берёт `consents.py`.

def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    made = [build_inspection(), build_reply(), build_log(), build_notice(),
            build_docs_request()]
    if "--pdf" in sys.argv:
        made += [pdf for pdf in (B.build_pdf(d) for d in list(made)) if pdf]
    for f in made:
        print(f"  ✓ {f.relative_to(HERE.parent.parent)}  "
              f"({f.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
